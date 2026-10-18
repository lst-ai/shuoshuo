"""把會議紀錄（Markdown）存成 Word 檔。只處理本工具會產生的簡單格式。"""
from __future__ import annotations

import re
from pathlib import Path


def _add_inline(paragraph, text: str) -> None:
    # **粗體**
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run = paragraph.add_run(part)
            run.bold = i % 2 == 1


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(md: str, path: str | Path) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微軟正黑體"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微軟正黑體")

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if _is_table_row(s):
            rows = []
            while i < len(lines) and _is_table_row(lines[i]):
                cells = _cells(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncol)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(ncol):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            # "#" → 文件標題，"##" → 標題1，"###" → 標題2
            doc.add_heading(m.group(2).replace("**", ""), level=len(m.group(1)) - 1)
        elif re.match(r"^[-*•]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.、)]\s*", s):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\d+[.、)]\s*", "", s))
        else:
            _add_inline(doc.add_paragraph(), s)
        i += 1
    doc.save(str(path))
