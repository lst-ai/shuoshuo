import json

import httpx
import numpy as np
import pytest

from shuoshuo import audio, config, polish, textproc, transcribe
from shuoshuo.export import markdown_to_docx
from shuoshuo.history import History


def test_traditional_and_replacements():
    out = textproc.postprocess("我们明天开会讨论软件 雲林線", {"雲林線": "雲林縣"})
    assert "我們" in out and "軟體" in out and "雲林縣" in out
    assert " " not in out


def test_replacement_longest_first():
    assert textproc.apply_replacements("AI方舟", {"方舟": "X", "AI方舟": "AI人才方舟計畫"}) == "AI人才方舟計畫"


def test_hallucination_removed():
    assert textproc.clean_transcript("今天開會。字幕由Amara.org社群提供") == "今天開會。"


def test_split_chunks_cuts_in_silence():
    sr = audio.SAMPLE_RATE
    loud = 0.3 * np.ones(sr * 595, dtype=np.float32)
    quiet = np.zeros(sr * 2, dtype=np.float32)
    x = np.concatenate([loud, quiet, loud])
    chunks = audio.split_chunks(x, max_sec=600)
    assert len(chunks) == 2
    assert sum(len(c) for c in chunks) == len(x)
    assert 595 * sr <= len(chunks[0]) <= 597 * sr  # 切在安靜處
    assert audio.to_flac_bytes(chunks[0][: sr * 5])[:4] == b"fLaC"


def test_silence_detect():
    assert audio.is_silent(np.zeros(16000, dtype=np.float32))
    assert not audio.is_silent(0.2 * np.ones(16000, dtype=np.float32))
    # 小聲說話：大部分時間安靜，只有一小段有聲音，不能被判為安靜
    quiet = np.zeros(16000 * 5, dtype=np.float32)
    quiet[16000:20000] = 0.01 * np.sin(np.linspace(0, 600, 4000))
    assert not audio.is_silent(quiet)


def test_silence_raises(keys):
    with pytest.raises(transcribe.SilenceError):
        transcribe.transcribe(np.zeros(32000, dtype=np.float32), config.Settings())


def test_history(tmp_path):
    h = History(tmp_path / "h.db")
    h.add(raw="原文", final="潤飾後的文字", style="日常", duration=3)
    h.add(kind="meeting", raw="x", final="# 會議紀錄")
    assert len(h.search()) == 2
    assert len(h.search("潤飾")) == 1
    assert h.recent_dictations()[0].final == "潤飾後的文字"
    h.delete(h.search("潤飾")[0].id)
    assert len(h.search()) == 1


def test_docx_export(tmp_path):
    md = """# 會議紀錄：課發會
## 決議事項
1. 通過**本土語**課程計畫
- 重點一
## 待辦事項
| 事項 | 負責人 | 期限 |
|---|---|---|
| 送審 | 教務處 | 10/31 |
"""
    p = tmp_path / "m.docx"
    markdown_to_docx(md, p)
    from docx import Document

    d = Document(str(p))
    assert d.tables[0].cell(1, 1).text == "教務處"
    assert any("本土語" in para.text for para in d.paragraphs)


@pytest.fixture
def keys(monkeypatch, tmp_path):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    monkeypatch.setenv("GROQ_API_KEY", "gsk-test")
    monkeypatch.setenv("APPDATA", str(tmp_path))
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path))


def test_polish_request(monkeypatch, keys):
    seen = {}

    def fake_post(url, headers, json, timeout):
        seen.update(url=url, headers=headers, body=json)
        return httpx.Response(200, json={"content": [{"type": "text", "text": "明天下午3點開會。"}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    s = config.Settings(vocabulary=["雲林縣"], polish_provider="claude")
    out = polish.polish("嗯明天那個下午三點開會", s, "正式公文")
    assert out == "明天下午3點開會。"
    assert seen["headers"]["x-api-key"] == "sk-test"
    assert seen["body"]["model"] == "claude-haiku-5-5"
    assert "正式公文" in seen["body"]["system"] and "雲林縣" in seen["body"]["system"]
    assert "<逐字稿>" in seen["body"]["messages"][0]["content"]


def test_polish_bad_key(monkeypatch, keys):
    monkeypatch.setattr(httpx, "post", lambda *a, **k: httpx.Response(401, json={}))
    with pytest.raises(polish.PolishError, match="無效"):
        polish.polish("測試", config.Settings(provider_fallback=False))


def test_groq_polish_default(monkeypatch, keys):
    seen = {}

    def fake_post(url, headers, json, timeout):
        seen.update(url=url, body=json, headers=headers)
        return httpx.Response(200, json={"choices": [{"message": {"content": "好的。", "reasoning": "想…"}}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    assert polish.polish("嗯好的", config.Settings()) == "好的。"
    assert "groq.com" in seen["url"] and seen["headers"]["Authorization"] == "Bearer gsk-test"
    assert seen["body"]["model"] == "openai/gpt-oss-120b"
    assert seen["body"]["include_reasoning"] is False
    assert seen["body"]["messages"][0]["role"] == "system"


def test_gemini_polish(monkeypatch, keys):
    monkeypatch.setenv("GEMINI_API_KEY", "gm-test")
    seen = {}

    def fake_post(url, headers, json, timeout):
        seen.update(url=url, body=json, headers=headers)
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [
            {"text": "思考中", "thought": True}, {"text": "會議紀錄"}]}}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    s = config.Settings(meeting_provider="gemini")
    assert polish.meeting_minutes("逐字稿", s) == "會議紀錄"
    assert "gemini-2.5-flash:generateContent" in seen["url"]
    assert seen["headers"]["x-goog-api-key"] == "gm-test"
    assert "會議紀錄整理員" in seen["body"]["systemInstruction"]["parts"][0]["text"]


def test_provider_fallback(monkeypatch, keys):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    urls = []

    def fake_post(url, headers, json, timeout):
        urls.append(url)
        if "groq.com" in url:
            return httpx.Response(429, json={"error": {"message": "rate limit"}})
        return httpx.Response(200, json={"content": [{"type": "text", "text": "Claude 救援"}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    assert polish.polish("測試", config.Settings()) == "Claude 救援"
    assert "groq.com" in urls[0] and "anthropic.com" in urls[-1]
    # Gemini 沒填金鑰 → 直接略過，不會呼叫
    assert not any("googleapis" in u for u in urls)
    # 會議預設 Gemini，沒金鑰時自動改用 Groq
    urls.clear()
    monkeypatch.setattr(httpx, "post", lambda url, headers, json, timeout: (urls.append(url), httpx.Response(
        200, json={"choices": [{"message": {"content": "紀錄"}}]}))[1])
    assert polish.meeting_minutes("逐字稿", config.Settings()) == "紀錄" and "groq.com" in urls[0]
    monkeypatch.setattr(httpx, "post", fake_post)
    # 關閉備援時就直接報錯，且訊息說明額度用完
    with pytest.raises(polish.PolishError, match="額度"):
        polish.polish("測試", config.Settings(provider_fallback=False))


def test_groq_and_fallback(monkeypatch, keys):
    calls = []

    def fake_post(url, headers, files, data, timeout):
        calls.append(data)
        return httpx.Response(200, json={"text": "我们开会"})

    monkeypatch.setattr(httpx, "post", fake_post)
    x = (0.2 * np.sin(np.linspace(0, 5000, 32000))).astype(np.float32)
    text, eng = transcribe.transcribe(x, config.Settings())
    assert text == "我們開會" and eng == "groq"
    assert calls[0]["language"] == "zh"

    # Groq 失敗 → 改用本機
    monkeypatch.setattr(httpx, "post", lambda *a, **k: httpx.Response(500, text="boom"))
    monkeypatch.setattr(transcribe, "transcribe_local", lambda x, s, p=None: "本机结果")
    text, eng = transcribe.transcribe(x, config.Settings())
    assert text == "本機結果" and eng == "local"


def test_settings_roundtrip(keys):
    s = config.Settings.load()
    s.vocabulary = ["AI人才方舟計畫"]
    s.custom_styles = {"校務會議": "條理分明"}
    s.save()
    s2 = config.Settings.load()
    assert s2.vocabulary == ["AI人才方舟計畫"] and s2.custom_styles["校務會議"] == "條理分明"
    raw = json.loads(config.Settings.path().read_text(encoding="utf-8"))
    assert "sk-test" not in json.dumps(raw)  # 金鑰不會寫進設定檔
